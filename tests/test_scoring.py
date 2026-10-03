from __future__ import annotations

from stellenradar import arbeitsagentur
from stellenradar.config import lade_config
from stellenradar.db import verbinden
from stellenradar.filter import Filter
from stellenradar.scoring import firmen_uebersicht, score_berechnen

from fake_ba import FakeSession


def test_score_formel():
    cfg = lade_config()
    cfg["scoring"].update(punkte_pro_stelle=10, max_punkte_stellen=50, punkte_pro_tag_offen=0.5,
                          max_punkte_alter=30, bonus_erneut_ausgeschrieben=20, bonus_regional=25,
                          abzug_personaldienstleister=30)
    assert score_berechnen(cfg, 3, 20, False, False, False)[0] == 40
    assert score_berechnen(cfg, 9, 200, True, True, False)[0] == 50 + 30 + 20 + 25
    assert score_berechnen(cfg, 1, 0, False, True, True)[0] == 10 + 25 - 30
    # Gewichtung aus der Config
    cfg["scoring"]["bonus_regional"] = 100
    assert score_berechnen(cfg, 1, 0, False, True, False)[0] == 110


def test_uebersicht(tmp_path):
    cfg = lade_config()
    cfg["arbeitsagentur"]["pause_sekunden"] = 0
    con = verbinden(tmp_path / "t.db")
    flt = Filter(cfg)
    arbeitsagentur.abrufen(con, cfg, flt, erzwingen=True,
                           client=arbeitsagentur.BAClient(cfg, FakeSession(), schlafen=lambda s: None))
    df = firmen_uebersicht(con, cfg, flt).set_index("firma_key")
    w = df.loc["weinheimer maschinenbau"]
    assert w["Offene Stellen"] == 3 and w["Regional"] and w["Älteste Anzeige (Tage)"] == 75
    b = df.loc["bergstrasse logistik"]
    assert b["Erneut ausgeschrieben"] and b["Älteste Anzeige (Tage)"] == 120   # Kette zählt
    assert df.loc["kurpfalz personalservice"]["Personaldienstleister?"]
    assert df.loc["sap"]["Blacklist"]
    assert df.loc["odenwald elektronik"]["Offene Stellen"] == 2
    # Sortierung nach Score
    assert list(df["Score"]) == sorted(df["Score"], reverse=True)
    # erster Abruf ist kein "Neu"
    assert (df["Neu"] == "").all()


def test_excel_export(tmp_path):
    import io

    import openpyxl

    from stellenradar.export import excel_bytes

    cfg = lade_config()
    cfg["arbeitsagentur"]["pause_sekunden"] = 0
    con = verbinden(tmp_path / "t.db")
    arbeitsagentur.abrufen(con, cfg, Filter(cfg), erzwingen=True,
                           client=arbeitsagentur.BAClient(cfg, FakeSession(), schlafen=lambda s: None))
    df = firmen_uebersicht(con, cfg)
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes(df, con)))
    assert wb.sheetnames == ["Firmen", "Anzeigen"]
    kopf = [c.value for c in wb["Firmen"][1]]
    assert "Firma" in kopf and "Score" in kopf and "LinkedIn" in kopf and "firma_key" not in kopf
    assert wb["Firmen"].max_row == len(df) + 1
    assert any(str(c.value).startswith("https://www.arbeitsagentur.de/jobsuche/jobdetail/")
               for c in wb["Anzeigen"]["L"][1:])
